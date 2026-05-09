"""
Gerador de slides PPTX para a apresentação da P2.

Uso:
    uv run python docs/build_slides.py

Produz docs/apresentacao-p2.pptx (~16 slides) com tema escuro e accent verde.
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


# ── Tema ────────────────────────────────────────────────────────────────
BG_DARK = RGBColor(0x0F, 0x17, 0x1F)
BG_PANEL = RGBColor(0x18, 0x22, 0x2E)
TEXT_PRIMARY = RGBColor(0xE6, 0xED, 0xF3)
TEXT_MUTED = RGBColor(0x9A, 0xA8, 0xB6)
ACCENT = RGBColor(0x3F, 0xB9, 0x50)
ACCENT_DIM = RGBColor(0x23, 0x6E, 0x35)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_background(slide, color: RGBColor) -> None:
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.shadow.inherit = False
    # Mover para o fundo
    spTree = bg._element.getparent()
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)


def add_text_box(
    slide,
    left: Emu,
    top: Emu,
    width: Emu,
    height: Emu,
    text: str,
    size: int = 18,
    bold: bool = False,
    color: RGBColor = TEXT_PRIMARY,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    font_name: str = "Helvetica Neue",
):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return tb


def add_bullets(slide, left: Emu, top: Emu, width: Emu, height: Emu, items: list[str], size: int = 16):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = f"•  {item}"
        run.font.size = Pt(size)
        run.font.color.rgb = TEXT_PRIMARY
        run.font.name = "Helvetica Neue"


def add_accent_bar(slide, top: Emu = Inches(0.45), width: Emu = Inches(0.18), height: Emu = Inches(0.5)):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), top, width, height)
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT


def add_footer(slide, slide_num: int, total: int):
    add_text_box(
        slide, Inches(0.55), Inches(7.05), Inches(6), Inches(0.3),
        "ICEV Remote Sensing Software · P2 Engenharia de Software",
        size=10, color=TEXT_MUTED,
    )
    add_text_box(
        slide, Inches(11.5), Inches(7.05), Inches(1.3), Inches(0.3),
        f"{slide_num} / {total}",
        size=10, color=TEXT_MUTED, align=PP_ALIGN.RIGHT,
    )


def make_slide(prs, layout_idx: int = 6):
    slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    add_background(slide, BG_DARK)
    return slide


def slide_title(slide, title: str, kicker: str = ""):
    add_accent_bar(slide)
    if kicker:
        add_text_box(
            slide, Inches(0.95), Inches(0.42), Inches(10), Inches(0.3),
            kicker.upper(), size=11, color=ACCENT, bold=True,
        )
    add_text_box(
        slide, Inches(0.95), Inches(0.7), Inches(11.5), Inches(0.7),
        title, size=30, bold=True,
    )


def add_panel(slide, left, top, width, height):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    panel.line.fill.background()
    panel.fill.solid()
    panel.fill.fore_color.rgb = BG_PANEL
    panel.adjustments[0] = 0.04
    return panel


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    total = 16  # número de slides; ajustar se mudar

    # ─ Slide 1 — Capa ────────────────────────────────────────────────────
    slide = make_slide(prs)
    add_accent_bar(slide, top=Inches(2.5), width=Inches(0.18), height=Inches(2.5))
    add_text_box(slide, Inches(0.95), Inches(2.35), Inches(11.5), Inches(0.4),
                 "AVALIAÇÃO P2 · ENGENHARIA DE SOFTWARE", size=14, color=ACCENT, bold=True)
    add_text_box(slide, Inches(0.95), Inches(2.85), Inches(11.5), Inches(1.2),
                 "ICEV Remote Sensing Software", size=44, bold=True)
    add_text_box(slide, Inches(0.95), Inches(4.0), Inches(11.5), Inches(0.6),
                 "Plataforma Web de Sensoriamento Remoto para Agricultura de Precisão",
                 size=20, color=TEXT_MUTED)
    add_text_box(slide, Inches(0.95), Inches(5.4), Inches(11.5), Inches(0.4),
                 "Equipe", size=12, color=ACCENT, bold=True)
    add_text_box(slide, Inches(0.95), Inches(5.75), Inches(11.5), Inches(0.6),
                 "Luis Gustavo Olimpio · Lauan Matheus · João Leonardi · João Vinícius Castello",
                 size=14, color=TEXT_PRIMARY)
    add_text_box(slide, Inches(0.95), Inches(6.05), Inches(11.5), Inches(0.6),
                 "Vinícius Henrique · Lucas Benevinuto · José Melquíades",
                 size=14, color=TEXT_PRIMARY)
    add_text_box(slide, Inches(0.95), Inches(6.7), Inches(11.5), Inches(0.4),
                 "Maio · 2026", size=12, color=TEXT_MUTED)

    # ─ Slide 2 — Agenda ──────────────────────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "Agenda", kicker="Roteiro de 20 min")
    add_bullets(slide, Inches(0.95), Inches(1.8), Inches(11.5), Inches(5),
                [
                    "Problema e contexto — agricultura de precisão no Brasil",
                    "Solução — visão geral do produto",
                    "Backlog priorizado e funcionalidades do MVP",
                    "Arquitetura — frontend, backend, dados",
                    "Tecnologias e decisões de design",
                    "Demonstração ao vivo (NDVI, 3D, classificação)",
                    "Estratégia de testes e resultados",
                    "Encerramento, próximos passos e dúvidas",
                ], size=18)
    add_footer(slide, 2, total)

    # ─ Slide 3 — Problema ────────────────────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "O problema", kicker="Por que existir?")
    add_text_box(slide, Inches(0.95), Inches(1.8), Inches(11.5), Inches(0.6),
                 "Monitorar saúde de culturas é caro, complexo e fragmentado.",
                 size=22, bold=True)
    add_panel(slide, Inches(0.95), Inches(2.7), Inches(11.5), Inches(3.6))
    add_bullets(slide, Inches(1.25), Inches(2.95), Inches(11), Inches(3.2),
                [
                    "Softwares GIS comerciais (ArcGIS, ENVI) custam milhares de R$ por ano",
                    "Ferramentas profissionais exigem treinamento especializado",
                    "Ausência de soluções integradas: visualização + análise + classificação",
                    "Imagens Sentinel-2 (gratuitas, 10 m, revisita 5 dias) são subutilizadas no Brasil",
                    "Decisões tardias = perdas de produtividade e custos operacionais",
                ], size=16)
    add_footer(slide, 3, total)

    # ─ Slide 4 — Solução ─────────────────────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "A solução", kicker="ICEV Remote Sensing Software")
    add_text_box(slide, Inches(0.95), Inches(1.7), Inches(11.5), Inches(1.2),
                 "Plataforma web local que processa imagens Sentinel-2 sobre talhões e gera índices, classificações e visualizações 3D — sem custo de licenciamento.",
                 size=18, color=TEXT_PRIMARY)
    # 4 caixas de feature
    coords = [(0.95, 3.2), (4.45, 3.2), (7.95, 3.2), (0.95, 5.0)]
    titles = ["Importar Talhões", "Calcular Índices", "Classificar Zonas", "Visualizar em 3D"]
    descs = [
        "KML do Google Earth ou desenho direto no mapa",
        "NDVI · EVI · SAVI · NDWI · NDBI + 13 bandas",
        "K-Means · Threshold · Supervisionada simulada",
        "Renderização Three.js com elevação por valor",
    ]
    for i, ((cx, cy), t, d) in enumerate(zip(coords, titles, descs)):
        add_panel(slide, Inches(cx), Inches(cy), Inches(3.4), Inches(1.7))
        add_text_box(slide, Inches(cx + 0.2), Inches(cy + 0.2), Inches(3.0), Inches(0.4),
                     t, size=15, bold=True, color=ACCENT)
        add_text_box(slide, Inches(cx + 0.2), Inches(cy + 0.65), Inches(3.0), Inches(1.0),
                     d, size=12, color=TEXT_MUTED)
    add_footer(slide, 4, total)

    # ─ Slide 5 — Usuários-alvo ────────────────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "Para quem é?", kicker="Usuários-alvo")
    personas = [
        ("Produtor Rural", "Monitorar lavouras, identificar áreas de estresse, otimizar insumos"),
        ("Agrônomo", "Mapas de variabilidade, prescrição de insumos, análise de vigor"),
        ("Pesquisador", "Experimentar algoritmos sobre dados reais, comparar índices"),
        ("Estudante", "Aprender sensoriamento remoto na prática"),
    ]
    for i, (name, desc) in enumerate(personas):
        y = 1.85 + i * 1.1
        add_panel(slide, Inches(0.95), Inches(y), Inches(11.5), Inches(1.0))
        add_text_box(slide, Inches(1.2), Inches(y + 0.15), Inches(3.5), Inches(0.6),
                     name, size=18, bold=True, color=ACCENT)
        add_text_box(slide, Inches(4.7), Inches(y + 0.18), Inches(7.5), Inches(0.7),
                     desc, size=14, color=TEXT_PRIMARY)
    add_footer(slide, 5, total)

    # ─ Slide 6 — Backlog ─────────────────────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "Backlog priorizado", kicker="6 épicos · 33 histórias")
    items = [
        "Épico 1 — Gestão de Talhões (KML, mapa, desenho)              · 6 histórias",
        "Épico 2 — Índices Espectrais (NDVI/EVI/SAVI/NDWI/NDBI/bandas) · 9 histórias",
        "Épico 3 — Visualização & Mapa Interativo (2D, 3D, popup)      · 5 histórias",
        "Épico 4 — Lab. de Experimentos (filtros, bordas, morfologia)  · 6 histórias",
        "Épico 5 — Classificação (K-Means, Threshold, Supervisionada)  · 6 histórias",
        "Épico 6 — Gestão de Dados Sentinel-2 (multi-produto)          · 1 história",
    ]
    add_bullets(slide, Inches(0.95), Inches(1.9), Inches(11.5), Inches(4),
                items, size=15)
    add_panel(slide, Inches(0.95), Inches(5.5), Inches(11.5), Inches(1.2))
    add_text_box(slide, Inches(1.2), Inches(5.7), Inches(11), Inches(0.4),
                 "MVP entregue: 28 das 33 histórias (84,8%)",
                 size=16, bold=True, color=ACCENT)
    add_text_box(slide, Inches(1.2), Inches(6.05), Inches(11), Inches(0.5),
                 "Pendentes: persistência de talhões desenhados · classificação supervisionada real · exportação de experimentos",
                 size=12, color=TEXT_MUTED)
    add_footer(slide, 6, total)

    # ─ Slide 7 — Arquitetura alto nível ──────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "Arquitetura", kicker="Visão geral")
    # caixa Frontend
    add_panel(slide, Inches(0.95), Inches(1.85), Inches(5.7), Inches(2.4))
    add_text_box(slide, Inches(1.15), Inches(2.0), Inches(5.3), Inches(0.4),
                 "FRONTEND · porta 3000", size=12, color=ACCENT, bold=True)
    add_text_box(slide, Inches(1.15), Inches(2.4), Inches(5.3), Inches(0.5),
                 "Next.js 16 · React 19 · TypeScript", size=15, bold=True)
    add_bullets(slide, Inches(1.15), Inches(2.95), Inches(5.3), Inches(1.3),
                ["Mapa Leaflet 2D + leaflet-draw",
                 "Three.js / @react-three/fiber (3D)",
                 "shadcn/ui + Tailwind CSS 4"], size=12)
    # caixa Backend
    add_panel(slide, Inches(6.85), Inches(1.85), Inches(5.7), Inches(2.4))
    add_text_box(slide, Inches(7.05), Inches(2.0), Inches(5.3), Inches(0.4),
                 "BACKEND · porta 8000", size=12, color=ACCENT, bold=True)
    add_text_box(slide, Inches(7.05), Inches(2.4), Inches(5.3), Inches(0.5),
                 "FastAPI + Uvicorn · Python 3.12", size=15, bold=True)
    add_bullets(slide, Inches(7.05), Inches(2.95), Inches(5.3), Inches(1.3),
                ["rasterio (JP2 Sentinel-2)",
                 "NumPy · SciPy · scikit-learn · OpenCV",
                 "Shapely · pyproj (geometria, reprojeção)"], size=12)
    # Dados
    add_panel(slide, Inches(0.95), Inches(4.5), Inches(11.6), Inches(1.8))
    add_text_box(slide, Inches(1.15), Inches(4.65), Inches(11), Inches(0.4),
                 "DADOS LOCAIS", size=12, color=ACCENT, bold=True)
    add_bullets(slide, Inches(1.15), Inches(5.05), Inches(11), Inches(1.2),
                [
                    "data/products/  →  Sentinel-2 .SAFE (granules em R10m, R20m, R60m)",
                    "data/KML Fields/  →  polígonos de talhões (KML do Google Earth)",
                ], size=13)
    add_text_box(slide, Inches(0.95), Inches(6.6), Inches(11.5), Inches(0.4),
                 "Comunicação REST/JSON · imagens em base64 · sem banco de dados",
                 size=12, color=TEXT_MUTED, align=PP_ALIGN.CENTER)
    add_footer(slide, 7, total)

    # ─ Slide 8 — Stack tecnológica ────────────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "Stack tecnológica", kicker="Por que escolhemos cada uma")
    rows = [
        ("FastAPI", "Validação Pydantic automática, OpenAPI grátis, async nativo"),
        ("rasterio", "Padrão de fato em I/O geoespacial Python (wrapper sobre GDAL)"),
        ("Next.js 16 + React 19", "App Router unifica páginas e rotas API server-side"),
        ("Leaflet", "Maduro, leve, integra com Esri/OSM sem token pago"),
        ("Three.js / R3F", "Visualização 3D nativa em WebGL, sem peso de Cesium"),
        ("Tailwind + shadcn/ui", "Componentes acessíveis (Radix), zero CSS-in-JS"),
        ("uv", "Build reprodutível e ~10× mais rápido que pip"),
    ]
    y = 1.85
    for tech, why in rows:
        add_panel(slide, Inches(0.95), Inches(y), Inches(11.5), Inches(0.65))
        add_text_box(slide, Inches(1.15), Inches(y + 0.12), Inches(3.5), Inches(0.45),
                     tech, size=14, bold=True, color=ACCENT)
        add_text_box(slide, Inches(4.55), Inches(y + 0.13), Inches(7.8), Inches(0.5),
                     why, size=12, color=TEXT_PRIMARY)
        y += 0.72
    add_footer(slide, 8, total)

    # ─ Slide 9 — Decisões de design ──────────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "Decisões de design", kicker="Trade-offs explícitos")
    decisions = [
        ("DA-01 · Frontend e backend em processos separados",
         "Permite que o processamento Python pesado não bloqueie a UI"),
        ("DA-02 · Imagens via base64 dentro do JSON",
         "API simples, uma resposta = todos os dados; trade-off: +33% de tamanho"),
        ("DA-04 · Reprojeção dinâmica (WGS84 → UTM)",
         "Frontend trabalha em lat/lon, backend trata projeções via pyproj"),
        ("DA-06 · SentinelProcessor isolado das rotas HTTP",
         "Permite testes unitários sem subir o servidor"),
        ("DA-07 · Tipagem end-to-end (TS + Pydantic)",
         "Erros de contrato detectados em build, não em produção"),
        ("DA-09 · Suavização gaussiana com máscara",
         "Evita 'sangramento' de valores nas bordas do polígono"),
    ]
    y = 1.8
    for title, desc in decisions:
        add_panel(slide, Inches(0.95), Inches(y), Inches(11.5), Inches(0.78))
        add_text_box(slide, Inches(1.15), Inches(y + 0.1), Inches(11), Inches(0.4),
                     title, size=13, bold=True, color=ACCENT)
        add_text_box(slide, Inches(1.15), Inches(y + 0.42), Inches(11), Inches(0.4),
                     desc, size=12, color=TEXT_MUTED)
        y += 0.85
    add_footer(slide, 9, total)

    # ─ Slide 10 — Fluxo NDVI ─────────────────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "Fluxo de cálculo de NDVI", kicker="Diagrama de sequência")
    steps = [
        ("1.", "Usuário seleciona talhão e índice 'NDVI' na UI"),
        ("2.", "Frontend envia POST /calculate-index com coordenadas WGS84"),
        ("3.", "Backend resolve produto Sentinel-2 .SAFE mais recente"),
        ("4.", "rasterio lê bandas B08 (NIR) e B04 (Red) em 10 m"),
        ("5.", "pyproj reprojeta polígono WGS84 → CRS nativo da imagem (UTM)"),
        ("6.", "rasterio.mask recorta as bandas ao polígono"),
        ("7.", "NumPy calcula NDVI = (NIR − Red) / (NIR + Red), clipa a [-1, 1]"),
        ("8.", "Pillow gera PNG colorido (gradiente marrom → amarelo → verde)"),
        ("9.", "Backend retorna JSON com base64, estatísticas, histograma e dados 3D"),
        ("10.", "Frontend sobrepõe imagem ao mapa Leaflet + atualiza painel de stats"),
    ]
    y = 1.85
    for n, txt in steps:
        add_text_box(slide, Inches(0.95), Inches(y), Inches(0.5), Inches(0.4),
                     n, size=14, bold=True, color=ACCENT)
        add_text_box(slide, Inches(1.5), Inches(y), Inches(11), Inches(0.5),
                     txt, size=14, color=TEXT_PRIMARY)
        y += 0.45
    add_footer(slide, 10, total)

    # ─ Slide 11 — Demo ───────────────────────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "Demonstração ao vivo", kicker="6 minutos")
    add_panel(slide, Inches(0.95), Inches(1.85), Inches(11.5), Inches(4.8))
    demo_steps = [
        "Abrir http://localhost:3000  →  selecionar talhão 'crop field 1'",
        "Calcular NDVI  →  ver overlay colorido + stats + histograma",
        "Clicar no mapa  →  popup com nível de estresse",
        "Trocar para EVI  →  comparar resultado",
        "Ativar visualização 3D  →  rotacionar terreno",
        "Ir em /classification  →  rodar K-Means com 3 classes  →  hectares por classe",
    ]
    y = 2.05
    for i, step in enumerate(demo_steps, start=1):
        add_text_box(slide, Inches(1.2), Inches(y), Inches(0.5), Inches(0.5),
                     f"{i}.", size=22, bold=True, color=ACCENT)
        add_text_box(slide, Inches(1.85), Inches(y + 0.05), Inches(10.4), Inches(0.6),
                     step, size=15, color=TEXT_PRIMARY)
        y += 0.7
    add_text_box(slide, Inches(0.95), Inches(6.85), Inches(11.5), Inches(0.4),
                 "Caso a internet falhe, prints da execução estão em docs/",
                 size=11, color=TEXT_MUTED, align=PP_ALIGN.CENTER)
    add_footer(slide, 11, total)

    # ─ Slide 12 — Estratégia de testes ───────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "Estratégia de testes", kicker="Plano em três níveis")
    add_panel(slide, Inches(0.95), Inches(1.85), Inches(3.7), Inches(2.7))
    add_text_box(slide, Inches(1.15), Inches(2.05), Inches(3.3), Inches(0.4),
                 "UNITÁRIO", size=12, bold=True, color=ACCENT)
    add_text_box(slide, Inches(1.15), Inches(2.45), Inches(3.3), Inches(0.5),
                 "pytest + NumPy", size=14, bold=True)
    add_bullets(slide, Inches(1.15), Inches(2.95), Inches(3.3), Inches(1.5),
                ["Cálculo de NDVI / EVI",
                 "Estatísticas",
                 "Histograma",
                 "Casos de borda"], size=11)
    add_panel(slide, Inches(4.85), Inches(1.85), Inches(3.7), Inches(2.7))
    add_text_box(slide, Inches(5.05), Inches(2.05), Inches(3.3), Inches(0.4),
                 "INTEGRAÇÃO", size=12, bold=True, color=ACCENT)
    add_text_box(slide, Inches(5.05), Inches(2.45), Inches(3.3), Inches(0.5),
                 "FastAPI TestClient", size=14, bold=True)
    add_bullets(slide, Inches(5.05), Inches(2.95), Inches(3.3), Inches(1.5),
                ["Endpoints /, /health, /products",
                 "Validação Pydantic (422)",
                 "Erros 400 / 404 / 500",
                 "Conversão coords → polígono"], size=11)
    add_panel(slide, Inches(8.75), Inches(1.85), Inches(3.7), Inches(2.7))
    add_text_box(slide, Inches(8.95), Inches(2.05), Inches(3.3), Inches(0.4),
                 "MANUAL", size=12, bold=True, color=ACCENT)
    add_text_box(slide, Inches(8.95), Inches(2.45), Inches(3.3), Inches(0.5),
                 "Demo end-to-end", size=14, bold=True)
    add_bullets(slide, Inches(8.95), Inches(2.95), Inches(3.3), Inches(1.5),
                ["Sentinel-2 real (1 GB)",
                 "Frontend Next.js",
                 "Renderização 3D",
                 "Apresentação ao vivo"], size=11)
    add_text_box(slide, Inches(0.95), Inches(4.85), Inches(11.5), Inches(0.4),
                 "Fixtures sintéticas (NumPy) substituem produtos reais — suite roda em < 1 s",
                 size=13, color=TEXT_MUTED, align=PP_ALIGN.CENTER)
    add_panel(slide, Inches(0.95), Inches(5.45), Inches(11.5), Inches(1.5))
    add_text_box(slide, Inches(1.2), Inches(5.65), Inches(11), Inches(0.4),
                 "Ver detalhes em docs/08-plano-de-testes.md", size=11, color=TEXT_MUTED)
    add_text_box(slide, Inches(1.2), Inches(6.0), Inches(11), Inches(0.5),
                 "$ uv run pytest tests/ -v --cov=api",
                 size=15, color=ACCENT, font_name="Menlo")
    add_text_box(slide, Inches(1.2), Inches(6.5), Inches(11), Inches(0.4),
                 "Comando único reproduz toda a suite", size=11, color=TEXT_MUTED)
    add_footer(slide, 12, total)

    # ─ Slide 13 — Resultados dos testes ──────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "Resultados dos testes", kicker="Execução de 09/05/2026")
    # 4 cartões com métricas
    metrics = [
        ("28", "casos executados"),
        ("100%", "taxa de sucesso"),
        ("0,79 s", "tempo total"),
        ("55%", "cobertura do core"),
    ]
    for i, (val, label) in enumerate(metrics):
        x = 0.95 + i * 2.95
        add_panel(slide, Inches(x), Inches(1.9), Inches(2.75), Inches(2.0))
        add_text_box(slide, Inches(x), Inches(2.15), Inches(2.75), Inches(1.0),
                     val, size=46, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
        add_text_box(slide, Inches(x), Inches(3.25), Inches(2.75), Inches(0.5),
                     label, size=13, color=TEXT_MUTED, align=PP_ALIGN.CENTER)
    add_panel(slide, Inches(0.95), Inches(4.2), Inches(11.5), Inches(2.65))
    add_text_box(slide, Inches(1.15), Inches(4.4), Inches(11), Inches(0.4),
                 "Cobertura por requisito", size=12, color=ACCENT, bold=True)
    add_bullets(slide, Inches(1.15), Inches(4.85), Inches(11), Inches(2),
                [
                    "RF-04 (produtos)  ·  RF-05 (NDVI)  ·  RF-06 (EVI)  ·  RF-10 (stats)",
                    "RF-11 (histograma)  ·  RF-20 (health)",
                    "RNF-04 (confiabilidade)  ·  RNF-05 (manutenibilidade)",
                    "Validação Pydantic exercitada com payloads inválidos (HTTP 422)",
                ], size=13)
    add_footer(slide, 13, total)

    # ─ Slide 14 — Riscos e limites ────────────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "Limitações conhecidas", kicker="O que não está no MVP")
    items = [
        ("Sem persistência de dados", "talhões desenhados são perdidos ao recarregar"),
        ("Classificação supervisionada simulada", "não usa modelo treinado real"),
        ("Imagens Sentinel-2 baixadas manualmente", "sem download automático via API Copernicus"),
        ("Single-user em localhost", "sem autenticação ou multi-tenancy"),
        ("Cobertura de I/O limitada", "read_band/crop dependem de .SAFE real → testes manuais"),
    ]
    y = 1.85
    for title, desc in items:
        add_panel(slide, Inches(0.95), Inches(y), Inches(11.5), Inches(0.85))
        add_text_box(slide, Inches(1.15), Inches(y + 0.1), Inches(11), Inches(0.4),
                     title, size=14, bold=True, color=ACCENT)
        add_text_box(slide, Inches(1.15), Inches(y + 0.45), Inches(11), Inches(0.4),
                     desc, size=12, color=TEXT_MUTED)
        y += 0.95
    add_footer(slide, 14, total)

    # ─ Slide 15 — Próximos passos ────────────────────────────────────────
    slide = make_slide(prs)
    slide_title(slide, "Próximos passos", kicker="Roadmap pós-MVP")
    add_bullets(slide, Inches(0.95), Inches(1.9), Inches(11.5), Inches(5),
                [
                    "Integração com a API do Copernicus (download automático)",
                    "Persistência via SQLite ou PostgreSQL/PostGIS",
                    "Modelo Random Forest treinado para classificação supervisionada real",
                    "Exportação dos resultados em GeoTIFF / Shapefile / CSV",
                    "Comparação temporal (multi-data) lado a lado",
                    "Aplicação web hospedada (multi-usuário, autenticação)",
                    "Pipeline CI com pytest + GitHub Actions",
                ], size=16)
    add_footer(slide, 15, total)

    # ─ Slide 16 — Encerramento ───────────────────────────────────────────
    slide = make_slide(prs)
    add_accent_bar(slide, top=Inches(2.8), width=Inches(0.18), height=Inches(2.0))
    add_text_box(slide, Inches(0.95), Inches(2.7), Inches(11.5), Inches(0.5),
                 "OBRIGADO", size=14, color=ACCENT, bold=True)
    add_text_box(slide, Inches(0.95), Inches(3.2), Inches(11.5), Inches(1.2),
                 "Perguntas e demonstração", size=44, bold=True)
    add_text_box(slide, Inches(0.95), Inches(4.4), Inches(11.5), Inches(0.6),
                 "github.com / icev-remote-sensing-software", size=16, color=TEXT_MUTED)
    add_text_box(slide, Inches(0.95), Inches(5.6), Inches(11.5), Inches(0.4),
                 "ENTREGÁVEIS DA P2", size=11, color=ACCENT, bold=True)
    add_text_box(slide, Inches(0.95), Inches(5.95), Inches(11.5), Inches(1.0),
                 "MVP funcional · Plano e relatório de testes · Manual do usuário · Documentação técnica · Apresentação",
                 size=14, color=TEXT_PRIMARY)

    out = Path(__file__).parent / "apresentacao-p2.pptx"
    prs.save(out)
    print(f"✓ Slides gerados em {out}")
    print(f"  Tamanho: {out.stat().st_size / 1024:.1f} KB")
    print(f"  Total de slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()
